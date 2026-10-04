from io import BytesIO
import openpyxl
from django.db import transaction
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from app_paths import database_path, temp_directory
from .models import Classroom, DeskCardTemplate, Meeting, MeetingParticipant, MeetingSeatAssignment, Participant, ParticipantCategory, PersonnelLevel, PoliceDepartment, Seat, SeatCellType

def _meeting_workspace_context():
    all_participants=Participant.objects.select_related('department','personnel_level').all()
    return {
        'meetings': Meeting.objects.select_related('venue').all()[:100],
        'venues': Classroom.objects.all().order_by('name'),
        'participants': all_participants.filter(active=True),
        'all_participants': all_participants,
        'levels': PersonnelLevel.objects.all(),
        'departments': PoliceDepartment.objects.all(),
        'desk_card_templates': {x.kind:x for x in DeskCardTemplate.objects.all()},
    }

def meeting_home(request):
    return render(request, 'seats/meeting_home.html', _meeting_workspace_context())

def participant_library(request):
    return render(request, 'seats/participant_library.html', _meeting_workspace_context())

def meeting_history(request):
    return render(request, 'seats/meeting_history.html', _meeting_workspace_context())

def meeting_settings(request):
    return render(request, 'seats/meeting_settings.html', _meeting_workspace_context())

@require_POST
def meeting_create(request):
    venue=get_object_or_404(Classroom, pk=request.POST.get('venue_id'))
    meeting=Meeting.objects.create(
        name=(request.POST.get('name') or '未命名会议').strip(),
        meeting_date=request.POST.get('meeting_date') or None,
        venue=venue,
        use_stage=request.POST.get('use_stage')=='on',
        stage_mode=request.POST.get('stage_mode') or 'specified',
    )
    ids=request.POST.getlist('participant_ids')
    for p in Participant.objects.filter(id__in=ids):
        MeetingParticipant.objects.create(meeting=meeting, participant=p)
    return redirect('meeting_detail', pk=meeting.pk)

def meeting_detail(request, pk):
    meeting=get_object_or_404(Meeting.objects.select_related('venue'), pk=pk)
    members=list(meeting.meeting_participants.select_related('participant__department','participant__personnel_level'))
    from .meeting_seating import participant_sort_key
    preview_members=sorted([m for m in members if m.include], key=lambda m: participant_sort_key(m.participant))
    preview=[m.participant for m in preview_members]
    assignments={a.seat_id:a for a in meeting.assignments.select_related('participant','seat')}
    grid=[]
    for r in range(1, meeting.venue.rows+1):
        row=[]
        for col in range(1, meeting.venue.cols+1):
            seat=next((s for s in meeting.venue.seats.all() if s.row==r and s.col==col),None)
            row.append((seat, assignments.get(seat.id) if seat else None))
        grid.append(row)
    return render(request,'seats/meeting_detail.html',{'meeting':meeting,'preview':preview,'preview_members':preview_members,'grid':grid})

@require_POST
def meeting_arrange(request, pk):
    meeting=get_object_or_404(Meeting, pk=pk)
    from .meeting_seating import auto_assign
    auto_assign(meeting)
    return redirect('meeting_detail', pk=pk)

@require_POST
def meeting_seat_state(request, pk, seat_id):
    meeting=get_object_or_404(Meeting, pk=pk)
    seat=get_object_or_404(Seat, pk=seat_id, classroom=meeting.venue)
    action=request.POST.get('action')
    assignment,_=MeetingSeatAssignment.objects.get_or_create(meeting=meeting,seat=seat)
    if action=='lock': assignment.locked=True
    elif action=='unlock': assignment.locked=False
    elif action=='skip': assignment.skipped=True; assignment.participant=None
    elif action=='unskip': assignment.skipped=False
    assignment.save()
    return redirect('meeting_detail',pk=pk)

@require_POST
def participant_create(request):
    Participant.objects.create(
        name=(request.POST.get('name') or '').strip(),
        category=request.POST.get('category') or ParticipantCategory.DEPARTMENT,
        department=PoliceDepartment.objects.filter(pk=request.POST.get('department_id') or None).first(),
        personnel_level=PersonnelLevel.objects.filter(pk=request.POST.get('level_id') or None).first(),
        position=(request.POST.get('position') or '').strip(),
        leader_order=int(request.POST.get('leader_order') or 100),
        attendee_order=int(request.POST.get('attendee_order') or 100),
        personal_order=int(request.POST.get('personal_order') or 100),
    )
    return redirect(request.POST.get('next') or 'meeting_home')

@require_POST
def department_create(request):
    PoliceDepartment.objects.create(name=request.POST['name'].strip(),short_name=(request.POST.get('short_name') or '').strip(),order=int(request.POST.get('order') or 100))
    return redirect(request.POST.get('next') or 'meeting_home')

@require_POST
def level_create(request):
    PersonnelLevel.objects.create(name=request.POST['name'].strip(),order=int(request.POST.get('order') or 100))
    return redirect(request.POST.get('next') or 'meeting_home')


@require_POST
def meeting_participant_state(request, pk, participant_id):
    meeting=get_object_or_404(Meeting, pk=pk)
    mp=get_object_or_404(MeetingParticipant, meeting=meeting, participant_id=participant_id)
    action=request.POST.get('action')
    if action=='stage' and meeting.use_stage:
        mp.is_stage=True
    elif action=='audience':
        mp.is_stage=False
    elif action=='exclude':
        mp.include=False
    elif action=='include':
        mp.include=True
    mp.save(update_fields=['is_stage','include'])
    return redirect('meeting_detail', pk=pk)

@require_POST
def meeting_swap_seats(request, pk):
    meeting=get_object_or_404(Meeting, pk=pk)
    seat_a=get_object_or_404(Seat, pk=request.POST.get('seat_a'), classroom=meeting.venue, cell_type=SeatCellType.SEAT)
    seat_b=get_object_or_404(Seat, pk=request.POST.get('seat_b'), classroom=meeting.venue, cell_type=SeatCellType.SEAT)
    if seat_a.pk == seat_b.pk: return JsonResponse({'ok':True})
    with transaction.atomic():
        a,_=MeetingSeatAssignment.objects.select_for_update().get_or_create(meeting=meeting, seat=seat_a)
        b,_=MeetingSeatAssignment.objects.select_for_update().get_or_create(meeting=meeting, seat=seat_b)
        if a.locked or b.locked or a.skipped or b.skipped:
            return JsonResponse({'ok':False,'error':'锁定或跳过座位不能交换'},status=400)
        pa,pb=a.participant,b.participant
        a.participant=None; a.save(update_fields=['participant'])
        b.participant=None; b.save(update_fields=['participant'])
        a.participant=pb; b.participant=pa
        a.save(update_fields=['participant']); b.save(update_fields=['participant'])
    return JsonResponse({'ok':True})

def meeting_print_cards(request, pk):
    meeting=get_object_or_404(Meeting, pk=pk)
    assigned=meeting.assignments.select_related('participant__department').exclude(participant=None).order_by('sort_order')
    defaults={'large':{'width_mm':190,'height_mm':90,'font_size_pt':52,'content_source':'person'},'small':{'width_mm':120,'height_mm':60,'font_size_pt':32,'content_source':'department'}}
    templates={x.kind:{'width_mm':x.width_mm,'height_mm':x.height_mm,'font_size_pt':x.font_size_pt,'content_source':x.content_source} for x in DeskCardTemplate.objects.filter(active=True)}
    for kind, values in defaults.items(): templates.setdefault(kind,values)
    cards=[]; seen_department_cards=set()
    for a in assigned:
        p=a.participant
        kind='small' if p.category == ParticipantCategory.DEPARTMENT else 'large'
        source=templates[kind]['content_source']
        if source == 'department' and p.department_id:
            dedupe=(kind,p.department_id)
            if dedupe in seen_department_cards: continue
            seen_department_cards.add(dedupe)
            text=p.department.desk_card_name
        else:
            text=p.name
        cards.append({'text':text,'kind':kind,'person':p})
    return render(request,'seats/meeting_cards.html',{'meeting':meeting,'cards':cards,'card_templates':templates})


@require_POST
def participant_update(request, participant_id):
    p=get_object_or_404(Participant, pk=participant_id)
    p.name=(request.POST.get('name') or p.name).strip()
    p.category=request.POST.get('category') or p.category
    p.department=PoliceDepartment.objects.filter(pk=request.POST.get('department_id') or None).first()
    p.personnel_level=PersonnelLevel.objects.filter(pk=request.POST.get('level_id') or None).first()
    p.position=(request.POST.get('position') or '').strip()
    p.leader_order=int(request.POST.get('leader_order') or 100)
    p.attendee_order=int(request.POST.get('attendee_order') or 100)
    p.personal_order=int(request.POST.get('personal_order') or 100)
    p.active=request.POST.get('active')=='on'
    p.remark=(request.POST.get('remark') or '').strip()
    p.save()
    return redirect(request.POST.get('next') or 'meeting_home')

@require_POST
def participant_toggle(request, participant_id):
    p=get_object_or_404(Participant, pk=participant_id)
    p.active=not p.active; p.save(update_fields=['active'])
    return redirect(request.POST.get('next') or 'meeting_home')

@require_POST
def meeting_update(request, pk):
    meeting=get_object_or_404(Meeting, pk=pk)
    meeting.name=(request.POST.get('name') or meeting.name).strip()
    meeting.meeting_date=request.POST.get('meeting_date') or None
    meeting.use_stage=request.POST.get('use_stage')=='on'
    meeting.stage_mode=request.POST.get('stage_mode') or 'specified'
    meeting.save()
    wanted={int(x) for x in request.POST.getlist('participant_ids') if x.isdigit()}
    existing={x.participant_id:x for x in meeting.meeting_participants.all()}
    for pid,mp in existing.items(): mp.include=pid in wanted; mp.save(update_fields=['include'])
    for pid in wanted-existing.keys():
        MeetingParticipant.objects.create(meeting=meeting,participant_id=pid)
    return redirect('meeting_detail',pk=pk)

@require_POST
def meeting_delete(request, pk):
    meeting=get_object_or_404(Meeting, pk=pk); meeting.delete()
    return redirect(request.POST.get('next') or 'meeting_home')

def participants_export(request):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='人员库'
    ws.append(['姓名','人员类别','行政级别','警种部门','职务','局领导顺序','列席顺序','部门内顺序','启用','备注'])
    for p in Participant.objects.select_related('department','personnel_level').all():
        ws.append([p.name,p.get_category_display(),p.personnel_level.name if p.personnel_level else '',p.department.name if p.department else '',p.position,p.leader_order,p.attendee_order,p.personal_order,'是' if p.active else '否',p.remark])
    out=BytesIO(); wb.save(out)
    resp=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition']="attachment; filename=participants.xlsx"
    return resp

@require_POST
def participants_import(request):
    f=request.FILES.get('file')
    if not f:
        return JsonResponse({'ok':False,'error':'请选择 Excel 文件'},status=400)

    def import_int(value, default=100):
        if value in (None, ''): return default
        try: return int(float(value))
        except (TypeError, ValueError): return default

    try:
        wb=openpyxl.load_workbook(f,data_only=True)
    except Exception:
        return JsonResponse({'ok':False,'error':'无法读取 Excel 文件，请确认文件格式正确'},status=400)
    ws=wb.active
    category_map={'局领导':ParticipantCategory.BUREAU_LEADER,'列席人员':ParticipantCategory.ATTENDEE,'警种部门':ParticipantCategory.DEPARTMENT,'警种部门人员':ParticipantCategory.DEPARTMENT}
    imported=0
    with transaction.atomic():
        for row in ws.iter_rows(min_row=2,values_only=True):
            if not row or not row[0]: continue
            vals=list(row)+[None]*10
            name=str(vals[0]).strip()
            level_name=str(vals[2]).strip() if vals[2] else ''
            dept_name=str(vals[3]).strip() if vals[3] else ''
            level=PersonnelLevel.objects.filter(name=level_name).first() if level_name else None
            dept=PoliceDepartment.objects.filter(name=dept_name).first() if dept_name else None
            # 同名人员可能属于不同部门；无部门人员则以“姓名+无部门”匹配。
            Participant.objects.update_or_create(name=name,department=dept,defaults={
                'category':category_map.get(str(vals[1]).strip(),ParticipantCategory.DEPARTMENT),
                'personnel_level':level,'position':str(vals[4] or '').strip(),
                'leader_order':import_int(vals[5]),'attendee_order':import_int(vals[6]),'personal_order':import_int(vals[7]),
                'active':str(vals[8] or '是').strip() not in ('否','0','False','false'),'remark':str(vals[9] or '').strip(),
            })
            imported += 1
    return redirect(request.POST.get('next') or 'meeting_home')


@require_POST
def desk_card_template_update(request):
    defaults={'large':(190,90,52,'person'),'small':(120,60,32,'department')}
    for kind,(dw,dh,df,ds) in defaults.items():
        obj,_=DeskCardTemplate.objects.get_or_create(kind=kind)
        def positive(name, fallback):
            try: return max(1,int(request.POST.get(f'{kind}_{name}') or fallback))
            except (TypeError,ValueError): return fallback
        obj.width_mm=positive('width',dw); obj.height_mm=positive('height',dh); obj.font_size_pt=positive('font_size',df)
        source=request.POST.get(f'{kind}_content_source') or ds
        obj.content_source=source if source in ('person','department') else ds
        obj.active=True; obj.save()
    return redirect(request.POST.get('next') or 'meeting_home')


def meeting_print_chart(request, pk):
    meeting=get_object_or_404(Meeting.objects.select_related('venue'), pk=pk)
    assignments={a.seat_id:a for a in meeting.assignments.select_related('participant__department','seat')}
    seats=list(meeting.venue.seats.all())
    grid=[]
    for r in range(1,meeting.venue.rows+1):
        row=[]
        for col in range(1,meeting.venue.cols+1):
            seat=next((s for s in seats if s.row==r and s.col==col),None)
            row.append((seat,assignments.get(seat.id) if seat else None))
        grid.append(row)
    paper=request.GET.get('paper','A4') if request.GET.get('paper') in ('A4','A3') else 'A4'
    orientation=request.GET.get('orientation','landscape')
    if orientation not in ('landscape','portrait'): orientation='landscape'
    return render(request,'seats/meeting_print_chart.html',{'meeting':meeting,'grid':grid,'paper':paper,'orientation':orientation})

def local_backup_download(request):
    from database_security import backup_database_for_update
    stamp=timezone.localtime().strftime('%Y%m%d-%H%M%S')
    dest=backup_database_for_update(f'manual-{stamp}')
    if not dest: raise Http404('数据库不存在')
    resp=HttpResponse(dest.read_bytes(),content_type='application/octet-stream')
    resp['Content-Disposition']=f'attachment; filename="meeting-seating-{stamp}.sqlite3"'
    return resp

@require_POST
def local_backup_restore(request):
    upload=request.FILES.get('file')
    if not upload: return JsonResponse({'ok':False,'error':'请选择备份文件'},status=400)
    if upload.size > 1024*1024*1024: return JsonResponse({'ok':False,'error':'备份文件过大'},status=400)
    target=database_path()
    if target.exists():
        from database_security import backup_database_for_update
        backup_database_for_update(f'before-restore-{timezone.localtime().strftime("%Y%m%d-%H%M%S")}')
    tmp=temp_directory()/'restore.sqlite3'
    with tmp.open('wb') as out:
        for chunk in upload.chunks(): out.write(chunk)
    from database_security import DatabaseSecurityError, verify_database_backup
    try:
        verify_database_backup(tmp)
    except DatabaseSecurityError:
        tmp.unlink(missing_ok=True)
        return JsonResponse({'ok':False,'error':'备份无效、已损坏，或不是由当前安装创建的加密备份'},status=400)
    from django.db import connections
    connections.close_all(); tmp.replace(target)
    return JsonResponse({'ok':True,'message':'恢复完成，请重新启动程序'})


def participant_edit(request, participant_id):
    p=get_object_or_404(Participant,pk=participant_id)
    return render(request,'seats/participant_edit.html',{'person':p,'levels':PersonnelLevel.objects.all(),'departments':PoliceDepartment.objects.all(),'category_choices':ParticipantCategory.choices})

def meeting_edit(request, pk):
    meeting=get_object_or_404(Meeting,pk=pk)
    selected=set(meeting.meeting_participants.filter(include=True).values_list('participant_id',flat=True))
    people=Participant.objects.select_related('department').filter(active=True)
    return render(request,'seats/meeting_edit.html',{'meeting':meeting,'participants':people,'selected':selected})

@require_POST
def department_update(request, department_id):
    d=get_object_or_404(PoliceDepartment,pk=department_id)
    d.name=(request.POST.get('name') or d.name).strip(); d.short_name=(request.POST.get('short_name') or '').strip()
    d.order=int(request.POST.get('order') or 100); d.active=request.POST.get('active')=='on'; d.save()
    return redirect(request.POST.get('next') or 'meeting_home')

@require_POST
def level_update(request, level_id):
    x=get_object_or_404(PersonnelLevel,pk=level_id)
    x.name=(request.POST.get('name') or x.name).strip(); x.order=int(request.POST.get('order') or 100)
    x.active=request.POST.get('active')=='on'; x.save()
    return redirect(request.POST.get('next') or 'meeting_home')
